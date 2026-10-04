"""Admin actions that protect legal records, and the accountant's reports."""
import io

import openpyxl
from conftest import PIN, set_schedule

from app import db

WORKDAYS = {str(d): "09:00-17:00/30" for d in range(6)}


def test_employee_with_real_punches_cannot_be_deleted(client, admin, clock, employee):
    with db.tx() as c:
        c.execute("INSERT INTO movements(employee_id, type, movement_at, created_at, mode, status, next_attempt_at) "
                  "VALUES (?,?,?,?,?,?,?)", (employee, "ARRIVAL", "2026-10-01T09:00:00", db.utc_now_iso(),
                                             "production", "submitted", db.utc_now_iso()))
    r = client.post(f"/admin/api/employees/{employee}/delete")
    assert r.status_code == 409
    assert db.one("SELECT COUNT(*) n FROM employees")["n"] == 1


def test_employee_with_only_test_punches_can_be_deleted(client, admin, clock, employee):
    with db.tx() as c:
        c.execute("INSERT INTO movements(employee_id, type, movement_at, created_at, mode, status, next_attempt_at) "
                  "VALUES (?,?,?,?,?,?,?)", (employee, "ARRIVAL", "2026-10-01T09:00:00", db.utc_now_iso(),
                                             "dry_run", "dry_run", db.utc_now_iso()))
    r = client.post(f"/admin/api/employees/{employee}/delete")
    assert r.status_code == 200
    assert db.one("SELECT COUNT(*) n FROM employees")["n"] == 0


def test_new_pin_rules(client, admin, employee):
    assert client.post(f"/admin/api/employees/{employee}/pin", json={"pin": "123456"}).status_code == 400
    assert client.post(f"/admin/api/employees/{employee}/pin", json={"pin": "12345"}).status_code == 400
    assert client.post(f"/admin/api/employees/{employee}/pin", json={"pin": "730194"}).status_code == 200


def test_purge_removes_only_test_punches(client, admin, clock, employee):
    with db.tx() as c:
        for mode in ("dry_run", "trial", "production"):
            c.execute("INSERT INTO movements(employee_id, type, movement_at, created_at, mode, status, next_attempt_at) "
                      "VALUES (?,?,?,?,?,?,?)", (employee, "ARRIVAL", "2026-10-01T09:00:00", db.utc_now_iso(),
                                                 mode, "pending", db.utc_now_iso()))
    assert client.post("/admin/api/movements/purge-tests").json()["deleted"] == 2
    assert [r["mode"] for r in db.all_rows("SELECT mode FROM movements")] == ["production"]


def _worked_week(client, employee, clock):
    """Mon 5 - Tue 6 Oct 2026: punches through the kiosk API (needs a registered device)."""
    from datetime import datetime
    set_schedule(client, employee, WORKDAYS)
    for day in (5, 6):
        for at, action in ((datetime(2026, 10, day, 9, 0), "ARRIVAL"), (datetime(2026, 10, day, 17, 5), "DEPARTURE")):
            clock.set(at)
            r = client.post("/api/kiosk/punch", json={"employee_id": employee, "pin": PIN, "action": action})
            assert r.status_code == 200, r.text
    clock.set(datetime(2026, 10, 7, 12, 0))


def test_monthly_report(kiosk, admin, clock, employee):
    _worked_week(kiosk, employee, clock)
    r = kiosk.get("/admin/api/report.xlsx?month=2026-10")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Σύνοψη", "Απολογιστικές δηλώσεις", "Αναλυτικά"]
    text = " ".join(str(c.value) for row in wb["Αναλυτικά"].iter_rows() for c in row if c.value is not None)
    assert "Παπαδοπούλου Μαρία" in text                            # official name, not the kiosk name
    assert "ΔΟΚΙΜΑΣΤΙΚΑ ΔΕΔΟΜΕΝΑ" in text                          # dry_run: marked as test data


def test_yearly_report(kiosk, admin, clock, employee):
    _worked_week(kiosk, employee, clock)
    r = kiosk.get("/admin/api/report-year.xlsx?year=2026")
    assert r.status_code == 200
    assert openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames == [
        "Σύνοψη", "Ανά μήνα", "Απολογιστικές δηλώσεις", "Αναλυτικά"]


def test_report_rejects_a_bad_month(client, admin):
    assert client.get("/admin/api/report.xlsx?month=2026-13").status_code == 400
