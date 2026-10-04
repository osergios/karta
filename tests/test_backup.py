"""Backups: the nightly status from backup.sh, its alerts, the database download and the yearly punch archive."""
import io
import json
import sqlite3
from datetime import datetime, timedelta, timezone

import openpyxl

from app import backupmark, config, db, monitor


def mark(local="ok", usb="-", cloud="-", hours_ago=0):
    assert backupmark.main([local, usb, cloud]) == 0
    if hours_ago:
        b = json.loads(db.setting("backup_status"))
        b["at"] = (datetime.now(timezone.utc) - timedelta(hours=hours_ago)).strftime("%Y-%m-%dT%H:%M:%S")
        with db.tx() as c:
            db.put_setting(c, "backup_status", json.dumps(b))


def alerts():
    return {a["kind"]: a for a in db.all_rows("SELECT kind, message FROM alerts")}


def test_backup_status_is_shown_and_ticks_the_first_step(client, admin):
    assert backupmark.main(["maybe", "-", "-"]) == 2
    mark(usb="ok")
    d = client.get("/admin/api/overview").json()
    assert d["backup"]["local"] == "ok" and d["backup"]["usb"] == "ok" and not d["backup"]["old"]
    assert {s["key"]: s for s in d["first_steps"]}["backup"]["done"]


def test_local_copy_only_is_not_enough_for_the_first_step(client, admin):
    mark()
    assert not {s["key"]: s for s in client.get("/admin/api/overview").json()["first_steps"]}["backup"]["done"]


def test_old_or_failed_backups_raise_an_alert():
    noon = datetime.now(config.TZ).replace(tzinfo=None, hour=12)
    mark(cloud="fail")
    monitor._check_backup(noon)
    assert "backup_failed" in alerts() and "cloud" in alerts()["backup_failed"]["message"]
    mark(hours_ago=60)
    monitor._check_backup(noon)
    assert "backup_old" in alerts()


def test_database_download_is_a_complete_copy(client, admin, employee):
    r = client.get("/admin/api/backup.db")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    path = config.DB_PATH + ".download-test"
    with open(path, "wb") as f:
        f.write(r.content)
    c = sqlite3.connect(path)
    assert c.execute("SELECT COUNT(*) FROM employees").fetchone()[0] == 1
    c.close()


def test_punch_archive_lists_every_real_punch_with_its_protocol(client, admin, employee):
    with db.tx() as c:
        for at, mode, status, proto in (("2026-03-02T09:01:00", "production", "submitted", "ΠΡΩΤ-123"),
                                        ("2026-03-02T17:00:00", "production", "local", None),
                                        ("2026-03-03T09:00:00", "dry_run", "dry_run", None),
                                        ("2025-12-31T09:00:00", "production", "submitted", "OLD")):
            c.execute("INSERT INTO movements(employee_id, type, movement_at, created_at, mode, status, protocol, "
                      "submitted_at, next_attempt_at) VALUES (?,?,?,?,?,?,?,?,?)",
                      (employee, "ARRIVAL", at, db.utc_now_iso(), mode, status, proto, "2026-03-02T07:01:05", db.utc_now_iso()))
    r = client.get("/admin/api/punches.xlsx?year=2026")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Χτυπήματα", "Δοκιμές"]
    rows = [row for row in wb["Χτυπήματα"].iter_rows(min_row=7, values_only=True) if row[0]]
    assert len(rows) == 2                                   # the 2025 punch and the test punch are not here
    assert rows[0][1] == "02/03/2026" and rows[0][8] == "ΠΡΩΤ-123" and rows[0][9] == "02/03/2026 09:01:05"
    assert rows[1][7] == "μόνο στην κάρτα"
    assert client.get("/admin/api/punches.xlsx?year=1999").status_code == 400
