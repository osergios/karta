"""«Αρχείο χτυπημάτων»: every punch of a year in one Excel file, readable without Karta.

The legal record of the digital work card has to be kept for years. Karta never deletes a real punch, and the
backups hold the whole database; this file is the plain copy for the archive and for an inspection: one row per
punch with its time, the employee, what happened to it and Ergani's protocol number.
"""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from . import config, db
from .report import LATE, STATUS, _employer_line, _info, _print_setup, _table_header, _widths

TYPE = {"ARRIVAL": "Προσέλευση", "DEPARTURE": "Αποχώρηση"}
HOW = {"pin": "PIN", "qr": "QR"}
COLS = ["Α/Α", "Ημερομηνία", "Ώρα", "Κίνηση", "Επώνυμο", "Όνομα", "ΑΦΜ εργαζόμενου", "Κατάσταση",
        "Αρ. πρωτοκόλλου ΕΡΓΑΝΗ", "Υποβλήθηκε στο ΕΡΓΑΝΗ", "Εκπρόθεσμη: αιτιολογία", "Τρόπος", "Συσκευή", "Σημείωση"]
WIDTHS = [7, 12, 9, 12, 20, 16, 13, 20, 18, 18, 30, 8, 16, 30]


def _local(utc_iso: str | None) -> str:
    if not utc_iso:
        return ""
    try:
        t = datetime.fromisoformat(utc_iso[:19]).replace(tzinfo=timezone.utc).astimezone(config.TZ)
    except ValueError:
        return utc_iso
    return f"{t:%d/%m/%Y %H:%M:%S}"


def _rows(year: int, real: bool):
    return db.all_rows(
        f"""SELECT m.*, e.last_name, e.first_name, e.afm, d.name AS device
            FROM movements m JOIN employees e ON e.id = m.employee_id LEFT JOIN devices d ON d.id = m.device_id
            WHERE m.movement_at >= ? AND m.movement_at < ? AND m.mode {'=' if real else '<>'} 'production'
            ORDER BY m.movement_at, m.id""", (f"{year:04d}-01-01", f"{year + 1:04d}-01-01"))


def _sheet(ws, year: int, rows, real: bool) -> None:
    _info(ws, [f"Αρχείο χτυπημάτων κάρτας εργασίας — {year}" + ("" if real else " — ΔΟΚΙΜΕΣ (χωρίς νομική ισχύ)"),
               _employer_line(),
               f"Δημιουργήθηκε {datetime.now(config.TZ):%d/%m/%Y %H:%M} από την Karta · {len(rows)} κινήσεις · "
               "ώρες Ελλάδας (Europe/Athens)",
               "Όλες οι κινήσεις όπως καταγράφηκαν. «Μόνο στην κάρτα» = καταχώριση διαχειριστή που δεν στάλθηκε στο "
               "ΕΡΓΑΝΗ (π.χ. ξεχασμένη αποχώρηση). Φυλάξτε το αρχείο μαζί με τα αντίγραφα ασφαλείας."],
          len(COLS), WIDTHS)
    head = 6
    _table_header(ws, head, COLS)
    for i, m in enumerate(rows, 1):
        at = datetime.fromisoformat(m["movement_at"])
        keys = m.keys()
        ws.append([i, at.strftime("%d/%m/%Y"), at.strftime("%H:%M:%S"), TYPE.get(m["type"], m["type"]),
                   m["last_name"], m["first_name"], m["afm"], STATUS.get(m["status"], m["status"]),
                   m["protocol"] or "", _local(m["submitted_at"]) if m["status"] == "submitted" else "",
                   LATE.get(m["late_justification"] or "", m["late_justification"] or ""),
                   HOW.get(m["auth_method"] if "auth_method" in keys else None, "διαχειριστής" if m["status"] == "local" else ""),
                   m["device"] or "", (m["note"] if "note" in keys else None) or ""])
    if not rows:
        ws.cell(head + 1, 1, "Καμία κίνηση για αυτό το έτος.").font = Font(italic=True, color="8A8A8A")
    _widths(ws, WIDTHS)
    _print_setup(ws, head)


def build(year: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Χτυπήματα"
    _sheet(ws, year, _rows(year, True), True)
    tests = _rows(year, False)
    if tests:
        _sheet(wb.create_sheet("Δοκιμές"), year, tests, False)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
